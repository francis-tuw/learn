# -*- coding: utf-8 -*-
"""
Excel输出工具 - 按照指定格式生成Excel输出结果
"""

import os
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Excel输出器 - 按照test输出示例.xlsx格式生成结果"""
    
    def __init__(self):
        pass
    
    def generate_excel(self, comparison_results, output_path):
        """生成Excel输出文件
        
        Args:
            comparison_results (dict): 比较结果数据
            output_path (str): 输出文件路径
        
        Returns:
            bool: 生成是否成功
        """
        try:
            # 创建工作簿
            wb = Workbook()
            
            # 创建工作表
            ws = wb.active
            ws.title = "SQL检查结果"
            
            # 设置表头
            self._set_header(ws)
            
            # 填充数据
            self._fill_data(ws, comparison_results)
            
            # 设置列宽
            self._set_column_widths(ws)
            
            # 保存文件
            wb.save(output_path)
            
            return True
        except Exception as e:
            print(f"生成Excel文件失败: {e}")
            return False
    
    def _set_header(self, ws):
        """设置表头"""
        # 表头数据
        headers = [
            "组别编号", "字段序号", "目标字段中文名", "目标字段英文名", 
            "取数方式", "源表", "源字段英文名", "源字段中文名", 
            "默认值", "转换逻辑", "SQL表达式", "是否一致", "备注"
        ]
        
        # 设置表头样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                      top=Side(style='thin'), bottom=Side(style='thin'))
        
        # 填充表头
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
    
    def _fill_data(self, ws, comparison_results):
        """填充数据"""
        # 准备数据
        data_rows = self._prepare_data_rows(comparison_results)
        
        # 填充数据
        for row_idx, row_data in enumerate(data_rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                # 设置边框
                border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                              top=Side(style='thin'), bottom=Side(style='thin'))
                cell.border = border
                # 设置对齐方式
                cell.alignment = Alignment(horizontal="left", vertical="center")
    
    def _prepare_data_rows(self, comparison_results):
        """准备数据行"""
        data_rows = []
        
        # 处理字段映射结果
        field_mappings = comparison_results.get('field_mappings', [])
        
        for mapping in field_mappings:
            row = [
                mapping.get('group_id', ''),
                mapping.get('field_seq', ''),
                mapping.get('target_field_cn', ''),
                mapping.get('target_field', ''),
                mapping.get('extract_method', ''),
                mapping.get('source_table_alias', ''),
                mapping.get('source_field', ''),
                mapping.get('source_field_cn', ''),
                mapping.get('default_value', ''),
                mapping.get('transformation_logic', ''),
                mapping.get('sql_expression', ''),
                mapping.get('is_consistent', ''),
                mapping.get('remark', '')
            ]
            data_rows.append(row)
        
        return data_rows
    
    def _set_column_widths(self, ws):
        """设置列宽"""
        # 列宽设置
        column_widths = {
            1: 12,  # 组别编号
            2: 10,  # 字段序号
            3: 15,  # 目标字段中文名
            4: 15,  # 目标字段英文名
            5: 12,  # 取数方式
            6: 15,  # 源表
            7: 15,  # 源字段英文名
            8: 15,  # 源字段中文名
            9: 12,  # 默认值
            10: 20,  # 转换逻辑
            11: 20,  # SQL表达式
            12: 10,  # 是否一致
            13: 20   # 备注
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
    
    def generate_from_task_result(self, task_id, task_result, output_path):
        """从任务结果生成Excel文件
        
        Args:
            task_id (str): 任务ID
            task_result (dict): 任务结果
            output_path (str): 输出文件路径
        
        Returns:
            bool: 生成是否成功
        """
        try:
            # 提取比较结果
            comparison_results = {
                'field_mappings': []
            }
            
            # 处理findings
            findings = task_result.get('findings', [])
            
            for finding in findings:
                field_mapping = {
                    'group_id': finding.get('group_id', ''),
                    'field_seq': finding.get('field_seq', ''),
                    'target_field_cn': finding.get('target_field_cn', ''),
                    'target_field': finding.get('target_field', ''),
                    'extract_method': finding.get('extract_method', ''),
                    'source_table_alias': finding.get('source_table_alias', ''),
                    'source_field': finding.get('source_field', ''),
                    'source_field_cn': finding.get('source_field_cn', ''),
                    'default_value': finding.get('default_value', ''),
                    'transformation_logic': finding.get('transformation_logic', ''),
                    'sql_expression': finding.get('original_sql', ''),
                    'is_consistent': '否' if finding.get('status') == 'inconsistent' else '是',
                    'remark': finding.get('explanation', '')
                }
                comparison_results['field_mappings'].append(field_mapping)
            
            # 生成Excel
            return self.generate_excel(comparison_results, output_path)
        except Exception as e:
            print(f"从任务结果生成Excel失败: {e}")
            return False


def export_to_excel(comparison_results, output_path):
    """导出到Excel的便捷函数
    
    Args:
        comparison_results (dict): 比较结果数据
        output_path (str): 输出文件路径
    
    Returns:
        bool: 导出是否成功
    """
    exporter = ExcelExporter()
    return exporter.generate_excel(comparison_results, output_path)


def export_from_task_result(task_id, task_result, output_dir):
    """从任务结果导出到Excel
    
    Args:
        task_id (str): 任务ID
        task_result (dict): 任务结果
        output_dir (str): 输出目录
    
    Returns:
        str: 生成的Excel文件路径
    """
    # 生成文件名
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(output_dir, f"SQL检查结果_{timestamp}.xlsx")
    
    # 确保输出目录存在
    os.makedirs(output_dir, exist_ok=True)
    
    # 生成Excel
    exporter = ExcelExporter()
    success = exporter.generate_from_task_result(task_id, task_result, output_path)
    
    if success:
        return output_path
    else:
        return None