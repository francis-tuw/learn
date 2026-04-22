#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
调试脚本，查看IT-Mapping文件中表的所属分组信息
"""

import openpyxl

def load_excel(file_path):
    """加载Excel文件"""
    return openpyxl.load_workbook(file_path, read_only=True, data_only=True)

def debug_table_groups(file_path):
    """调试表的所属分组信息"""
    print("=" * 60)
    print("调试表的所属分组信息")
    print("=" * 60)
    
    workbook = load_excel(file_path)
    
    # 找到数据源sheet
    sheet = None
    for name in workbook.sheetnames:
        if '数据源' in name or 'DataSet' in name:
            sheet = workbook[name]
            break
    
    if not sheet:
        print("未找到数据源sheet")
        workbook.close()
        return
    
    print(f"数据源sheet: {sheet.title}")
    print("\n表信息（前20行）:")
    print("-" * 100)
    print(f"{'行号':<5} {'序号':<5} {'所属系统':<15} {'数据表名_英文':<40} {'数据表别名':<10} {'数据表名_中文':<20} {'说明':<30} {'分组标识':<10}")
    print("-" * 100)
    
    for row_idx, row in enumerate(sheet.iter_rows(min_row=3, max_row=20, values_only=True), 3):
        if row[0] is not None and isinstance(row[0], int):
            # 提取信息
            seq = row[0]
            system = row[1] or ''
            table_en = row[2] or ''
            alias = row[3] or ''
            table_cn = row[4] or ''
            description = row[5] if len(row) > 5 else ''
            
            # 提取分组标识
            group_id = ''
            if description:
                if 'MP1' in str(description):
                    group_id = 'MP1'
                elif 'MP2' in str(description):
                    group_id = 'MP2'
                elif 'MP3' in str(description):
                    group_id = 'MP3'
            
            print(f"{row_idx:<5} {seq:<5} {system:<15} {table_en:<40} {alias:<10} {table_cn:<20} {str(description):<30} {group_id:<10}")
    
    print("-" * 100)
    
    # 检查分组关联信息
    print("\n分组关联信息:")
    print("-" * 100)
    print(f"{'行号':<5} {'组别编号':<10} {'组别名称':<15} {'表间关联关系及筛选条件':<50}")
    print("-" * 100)
    
    for row_idx, row in enumerate(sheet.iter_rows(min_row=14, max_row=30, values_only=True), 14):
        if row[0] is not None and row[0] != '组别编号':
            group_id = row[0]
            group_name = row[1] or ''
            relations = row[3] if len(row) > 3 else ''
            print(f"{row_idx:<5} {group_id:<10} {group_name:<15} {str(relations)[:50]:<50}")
    
    print("-" * 100)
    workbook.close()

if __name__ == '__main__':
    test_file = r'c:\Users\14093\Desktop\SQL智能纠错系统\【客户物理地址信息】ITMappingv0.8.xlsx'
    debug_table_groups(test_file)