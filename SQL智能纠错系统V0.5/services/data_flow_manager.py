# -*- coding: utf-8 -*-
"""
数据流处理模块
实现IT-Mapping文件和SQL存储过程的上传、解析、存储和一致性对比的完整流程
"""

import os
import re
import json
import uuid
import time
import threading
import hashlib
import logging
from datetime import datetime
from dataclasses import dataclass, field, asdict
from typing import Dict, List, Any, Optional, Tuple, Callable, Iterator
from enum import Enum
from concurrent.futures import ThreadPoolExecutor, as_completed
from functools import wraps
import traceback

import openpyxl
from openpyxl import load_workbook


class DataFlowPhase(Enum):
    UPLOAD = "upload"
    PARSE = "parse"
    STORE = "store"
    COMPARE = "compare"
    COMPLETE = "complete"
    ERROR = "error"


class DataFlowStatus(Enum):
    PENDING = "pending"
    IN_PROGRESS = "in_progress"
    COMPLETED = "completed"
    FAILED = "failed"
    RETRY = "retry"


@dataclass
class ProgressInfo:
    current_phase: str = ""
    current_step: str = ""
    total_items: int = 0
    processed_items: int = 0
    failed_items: int = 0
    percentage: float = 0.0
    start_time: Optional[datetime] = None
    estimated_remaining: int = 0
    message: str = ""
    
    def to_dict(self) -> Dict:
        return {
            'current_phase': self.current_phase,
            'current_step': self.current_step,
            'total_items': self.total_items,
            'processed_items': self.processed_items,
            'failed_items': self.failed_items,
            'percentage': round(self.percentage, 2),
            'start_time': self.start_time.isoformat() if self.start_time else None,
            'estimated_remaining': self.estimated_remaining,
            'message': self.message
        }


@dataclass
class ErrorInfo:
    error_type: str
    error_message: str
    error_phase: str
    error_step: str
    timestamp: datetime
    stack_trace: Optional[str] = None
    retry_count: int = 0
    is_recoverable: bool = True
    
    def to_dict(self) -> Dict:
        return {
            'error_type': self.error_type,
            'error_message': self.error_message,
            'error_phase': self.error_phase,
            'error_step': self.error_step,
            'timestamp': self.timestamp.isoformat(),
            'stack_trace': self.stack_trace,
            'retry_count': self.retry_count,
            'is_recoverable': self.is_recoverable
        }


@dataclass
class ValidationResult:
    is_valid: bool
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    file_hash: str = ""
    file_size: int = 0
    
    def to_dict(self) -> Dict:
        return {
            'is_valid': self.is_valid,
            'errors': self.errors,
            'warnings': self.warnings,
            'file_hash': self.file_hash,
            'file_size': self.file_size
        }


@dataclass
class ComparisonMetrics:
    total_records: int = 0
    consistent_records: int = 0
    inconsistent_records: int = 0
    warning_records: int = 0
    error_records: int = 0
    skipped_records: int = 0
    comparison_time: float = 0.0
    avg_comparison_time: float = 0.0
    
    def to_dict(self) -> Dict:
        return {
            'total_records': self.total_records,
            'consistent_records': self.consistent_records,
            'inconsistent_records': self.inconsistent_records,
            'warning_records': self.warning_records,
            'error_records': self.error_records,
            'skipped_records': self.skipped_records,
            'comparison_time': round(self.comparison_time, 2),
            'avg_comparison_time': round(self.avg_comparison_time, 4)
        }


class RetryMechanism:
    """重试机制 - 支持指数退避和自定义重试策略"""
    
    def __init__(self, max_retries: int = 3, base_delay: float = 1.0, 
                 max_delay: float = 60.0, exponential_base: float = 2.0):
        self.max_retries = max_retries
        self.base_delay = base_delay
        self.max_delay = max_delay
        self.exponential_base = exponential_base
        self.logger = logging.getLogger(__name__)
    
    def calculate_delay(self, retry_count: int) -> float:
        delay = self.base_delay * (self.exponential_base ** retry_count)
        return min(delay, self.max_delay)
    
    def execute_with_retry(self, func: Callable, *args, 
                          retryable_exceptions: Tuple = (Exception,),
                          on_retry: Optional[Callable] = None, **kwargs) -> Any:
        last_exception = None
        
        for retry_count in range(self.max_retries + 1):
            try:
                return func(*args, **kwargs)
            except retryable_exceptions as e:
                last_exception = e
                
                if retry_count < self.max_retries:
                    delay = self.calculate_delay(retry_count)
                    self.logger.warning(
                        f"操作失败，{delay:.1f}秒后重试 (第{retry_count + 1}/{self.max_retries}次): {str(e)}"
                    )
                    
                    if on_retry:
                        on_retry(retry_count, e, delay)
                    
                    time.sleep(delay)
                else:
                    self.logger.error(f"重试{self.max_retries}次后仍然失败: {str(e)}")
                    raise
        
        raise last_exception


class ComparisonLogger:
    """对比日志记录器 - 详细的日志记录和指标追踪"""
    
    def __init__(self, log_dir: str = "logs/comparison"):
        self.log_dir = log_dir
        os.makedirs(log_dir, exist_ok=True)
        self.logger = self._setup_logger()
        self.metrics = ComparisonMetrics()
        self.comparison_log: List[Dict] = []
        self.start_time: Optional[datetime] = None
    
    def _setup_logger(self) -> logging.Logger:
        logger = logging.getLogger('comparison_logger')
        logger.setLevel(logging.DEBUG)
        
        if not logger.handlers:
            fh = logging.FileHandler(
                os.path.join(self.log_dir, f'comparison_{datetime.now().strftime("%Y%m%d")}.log'),
                encoding='utf-8'
            )
            fh.setLevel(logging.DEBUG)
            
            formatter = logging.Formatter(
                '%(asctime)s - %(levelname)s - %(message)s',
                datefmt='%Y-%m-%d %H:%M:%S'
            )
            fh.setFormatter(formatter)
            logger.addHandler(fh)
        
        return logger
    
    def start_comparison(self, total_records: int):
        self.start_time = datetime.now()
        self.metrics.total_records = total_records
        self.logger.info(f"开始对比，总记录数: {total_records}")
    
    def log_record_comparison(self, record_id: str, status: str, 
                              details: Dict, duration: float = 0):
        log_entry = {
            'timestamp': datetime.now().isoformat(),
            'record_id': record_id,
            'status': status,
            'details': details,
            'duration_ms': round(duration * 1000, 2)
        }
        self.comparison_log.append(log_entry)
        
        if status == 'consistent':
            self.metrics.consistent_records += 1
            self.logger.debug(f"记录 {record_id}: 一致 - {details}")
        elif status == 'inconsistent':
            self.metrics.inconsistent_records += 1
            self.logger.warning(f"记录 {record_id}: 不一致 - {details}")
        elif status == 'warning':
            self.metrics.warning_records += 1
            self.logger.info(f"记录 {record_id}: 警告 - {details}")
        elif status == 'error':
            self.metrics.error_records += 1
            self.logger.error(f"记录 {record_id}: 错误 - {details}")
        elif status == 'skipped':
            self.metrics.skipped_records += 1
            self.logger.info(f"记录 {record_id}: 跳过 - {details}")
    
    def log_discrepancy(self, record_id: str, field_name: str,
                        mapping_value: Any, sql_value: Any, 
                        discrepancy_type: str, severity: str = 'medium'):
        discrepancy = {
            'record_id': record_id,
            'field_name': field_name,
            'mapping_value': str(mapping_value)[:500],
            'sql_value': str(sql_value)[:500],
            'discrepancy_type': discrepancy_type,
            'severity': severity
        }
        self.logger.warning(
            f"发现差异 [{severity}]: 记录={record_id}, 字段={field_name}, "
            f"类型={discrepancy_type}, Mapping值={mapping_value}, SQL值={sql_value}"
        )
        return discrepancy
    
    def end_comparison(self) -> ComparisonMetrics:
        if self.start_time:
            self.metrics.comparison_time = (datetime.now() - self.start_time).total_seconds()
        
        if self.metrics.total_records > 0:
            self.metrics.avg_comparison_time = (
                self.metrics.comparison_time / self.metrics.total_records
            )
        
        self.logger.info(
            f"对比完成: 总计={self.metrics.total_records}, "
            f"一致={self.metrics.consistent_records}, "
            f"不一致={self.metrics.inconsistent_records}, "
            f"警告={self.metrics.warning_records}, "
            f"错误={self.metrics.error_records}, "
            f"跳过={self.metrics.skipped_records}, "
            f"耗时={self.metrics.comparison_time:.2f}秒"
        )
        
        return self.metrics
    
    def export_log(self, filepath: str) -> bool:
        try:
            with open(filepath, 'w', encoding='utf-8') as f:
                json.dump({
                    'metrics': self.metrics.to_dict(),
                    'comparison_log': self.comparison_log
                }, f, ensure_ascii=False, indent=2)
            return True
        except Exception as e:
            self.logger.error(f"导出日志失败: {e}")
            return False


class DataUploadHandler:
    """数据上传处理器 - 处理文件上传、验证和错误处理"""
    
    ALLOWED_MAPPING_EXTENSIONS = {'.xlsx', '.xls'}
    ALLOWED_SQL_EXTENSIONS = {'.sql', '.txt'}
    MAX_FILE_SIZE = 100 * 1024 * 1024  # 100MB
    
    def __init__(self, upload_dir: str, logger: Optional[logging.Logger] = None):
        self.upload_dir = upload_dir
        os.makedirs(upload_dir, exist_ok=True)
        self.logger = logger or logging.getLogger(__name__)
        self.retry_mechanism = RetryMechanism(max_retries=3)
    
    def validate_file(self, file_path: str, file_type: str) -> ValidationResult:
        errors = []
        warnings = []
        
        if not os.path.exists(file_path):
            errors.append(f"文件不存在: {file_path}")
            return ValidationResult(is_valid=False, errors=errors)
        
        file_size = os.path.getsize(file_path)
        if file_size == 0:
            errors.append("文件为空")
        
        if file_size > self.MAX_FILE_SIZE:
            errors.append(f"文件大小({file_size / 1024 / 1024:.2f}MB)超过限制({self.MAX_FILE_SIZE / 1024 / 1024}MB)")
        
        ext = os.path.splitext(file_path)[1].lower()
        if file_type == 'mapping':
            if ext not in self.ALLOWED_MAPPING_EXTENSIONS:
                errors.append(f"不支持的文件格式: {ext}，支持的格式: {self.ALLOWED_MAPPING_EXTENSIONS}")
        elif file_type == 'sql':
            if ext not in self.ALLOWED_SQL_EXTENSIONS:
                errors.append(f"不支持的文件格式: {ext}，支持的格式: {self.ALLOWED_SQL_EXTENSIONS}")
        
        file_hash = self._calculate_file_hash(file_path)
        
        if file_type == 'mapping' and ext in self.ALLOWED_MAPPING_EXTENSIONS:
            try:
                wb = load_workbook(file_path, read_only=True)
                if len(wb.sheetnames) == 0:
                    warnings.append("Excel文件不包含任何工作表")
                wb.close()
            except Exception as e:
                errors.append(f"Excel文件损坏或格式不正确: {str(e)}")
        
        return ValidationResult(
            is_valid=len(errors) == 0,
            errors=errors,
            warnings=warnings,
            file_hash=file_hash,
            file_size=file_size
        )
    
    def _calculate_file_hash(self, file_path: str) -> str:
        hash_md5 = hashlib.md5()
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(4096), b''):
                hash_md5.update(chunk)
        return hash_md5.hexdigest()
    
    def save_uploaded_file(self, file_content: bytes, filename: str, 
                          file_type: str) -> Tuple[bool, str, Optional[ErrorInfo]]:
        try:
            unique_id = str(uuid.uuid4())[:8]
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_filename = re.sub(r'[^\w\-_\.]', '_', filename)
            new_filename = f"{timestamp}_{unique_id}_{safe_filename}"
            file_path = os.path.join(self.upload_dir, new_filename)
            
            def _write_file():
                with open(file_path, 'wb') as f:
                    f.write(file_content)
                return file_path
            
            saved_path = self.retry_mechanism.execute_with_retry(
                _write_file,
                retryable_exceptions=(IOError, OSError)
            )
            
            validation = self.validate_file(saved_path, file_type)
            if not validation.is_valid:
                return False, "", ErrorInfo(
                    error_type="ValidationError",
                    error_message="; ".join(validation.errors),
                    error_phase="upload",
                    error_step="validation",
                    timestamp=datetime.now(),
                    is_recoverable=False
                )
            
            self.logger.info(f"文件上传成功: {saved_path}")
            return True, saved_path, None
            
        except Exception as e:
            error_info = ErrorInfo(
                error_type=type(e).__name__,
                error_message=str(e),
                error_phase="upload",
                error_step="save",
                timestamp=datetime.now(),
                stack_trace=traceback.format_exc(),
                is_recoverable=True
            )
            self.logger.error(f"文件上传失败: {e}")
            return False, "", error_info
    
    def cleanup_old_files(self, max_age_days: int = 30) -> int:
        cleaned = 0
        cutoff = datetime.now().timestamp() - (max_age_days * 24 * 60 * 60)
        
        for filename in os.listdir(self.upload_dir):
            file_path = os.path.join(self.upload_dir, filename)
            if os.path.isfile(file_path):
                if os.path.getmtime(file_path) < cutoff:
                    try:
                        os.remove(file_path)
                        cleaned += 1
                    except Exception as e:
                        self.logger.warning(f"清理文件失败: {file_path}, {e}")
        
        return cleaned


class ITMappingProcessor:
    """IT-Mapping处理器 - 解析IT-Mapping文件并提取数据"""
    
    def __init__(self, progress_callback: Optional[Callable] = None):
        self.progress_callback = progress_callback
        self.logger = logging.getLogger(__name__)
        self.retry_mechanism = RetryMechanism(max_retries=2)
    
    def parse_mapping_file(self, file_path: str, 
                          task_id: str) -> Tuple[bool, Dict, Optional[ErrorInfo]]:
        try:
            self._update_progress(task_id, "parse", "加载Excel文件", 0, 100)
            
            def _load_workbook():
                return load_workbook(file_path, read_only=True, data_only=True)
            
            workbook = self.retry_mechanism.execute_with_retry(
                _load_workbook,
                retryable_exceptions=(Exception,)
            )
            
            file_hash = ""
            with open(file_path, 'rb') as f:
                file_hash = hashlib.md5(f.read()).hexdigest()
            
            result = {
                'file_path': file_path,
                'file_hash': file_hash,
                'sheets': {},
                'data_sources': [],
                'field_mappings': [],
                'relationships': [],
                'parse_metadata': {
                    'parse_time': datetime.now().isoformat(),
                    'total_sheets': len(workbook.sheetnames)
                }
            }
            
            total_sheets = len(workbook.sheetnames)
            
            for idx, sheet_name in enumerate(workbook.sheetnames):
                self._update_progress(
                    task_id, "parse", f"解析工作表: {sheet_name}",
                    idx + 1, total_sheets
                )
                
                sheet_data = self._parse_sheet(workbook[sheet_name], sheet_name)
                result['sheets'][sheet_name] = sheet_data
                
                if '数据源' in sheet_name or 'DataSet' in sheet_name:
                    result['data_sources'].extend(sheet_data.get('data_sources', []))
                elif '数据映射' in sheet_name or 'DataMap' in sheet_name:
                    result['field_mappings'].extend(sheet_data.get('field_mappings', []))
            
            workbook.close()
            
            self.logger.info(f"IT-Mapping解析完成: {file_path}, "
                           f"数据源={len(result['data_sources'])}, "
                           f"字段映射={len(result['field_mappings'])}")
            
            return True, result, None
            
        except Exception as e:
            error_info = ErrorInfo(
                error_type=type(e).__name__,
                error_message=str(e),
                error_phase="parse",
                error_step="mapping_file",
                timestamp=datetime.now(),
                stack_trace=traceback.format_exc(),
                is_recoverable=True
            )
            self.logger.error(f"IT-Mapping解析失败: {e}")
            return False, {}, error_info
    
    def _parse_sheet(self, sheet, sheet_name: str) -> Dict:
        result = {
            'sheet_name': sheet_name,
            'data_sources': [],
            'field_mappings': [],
            'row_count': 0
        }
        
        rows = list(sheet.iter_rows(values_only=True))
        result['row_count'] = len(rows)
        
        if len(rows) < 3:
            return result
        
        header_row = None
        for idx in range(min(5, len(rows))):
            row_values = [str(v).strip() if v else '' for v in rows[idx]]
            if any(keyword in ' '.join(row_values) for keyword in 
                   ['字段英文名', '目标字段', '组别编号', '数据表名']):
                header_row = idx
                break
        
        if header_row is None:
            return result
        
        headers = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(rows[header_row])]
        
        if '数据源' in sheet_name or 'DataSet' in sheet_name:
            result['data_sources'] = self._extract_data_sources(rows, headers, header_row)
        elif '数据映射' in sheet_name or 'DataMap' in sheet_name:
            result['field_mappings'] = self._extract_field_mappings(rows, headers, header_row)
        
        return result
    
    def _extract_data_sources(self, rows: List, headers: List[str], 
                             header_row: int) -> List[Dict]:
        data_sources = []
        
        for row_idx, row in enumerate(rows[header_row + 1:], header_row + 1):
            if not row or row[0] is None:
                continue
            
            row_dict = {headers[i]: row[i] if i < len(row) else None 
                       for i in range(len(headers))}
            
            if row_dict.get('组别编号') or row_dict.get('数据表名'):
                data_sources.append({
                    'row_index': row_idx,
                    'group_id': str(row_dict.get('组别编号', '')),
                    'table_name': str(row_dict.get('数据表名', '')),
                    'table_alias': str(row_dict.get('数据表别名', '')),
                    'table_name_cn': str(row_dict.get('数据表名(中文)', '')),
                    'source_system': str(row_dict.get('所属系统', '')),
                    'raw_data': row_dict
                })
        
        return data_sources
    
    def _extract_field_mappings(self, rows: List, headers: List[str], 
                               header_row: int) -> List[Dict]:
        field_mappings = []
        
        for row_idx, row in enumerate(rows[header_row + 1:], header_row + 1):
            if not row or row[0] is None:
                continue
            
            row_dict = {headers[i]: row[i] if i < len(row) else None 
                       for i in range(len(headers))}
            
            target_field = row_dict.get('字段英文名') or row_dict.get('目标字段')
            if target_field and str(target_field).strip():
                # 提取源表和源字段信息
                source_table = ''
                source_field = ''
                source_field_cn = ''
                
                # 尝试从不同的列名中提取信息
                for key in ['源表', '数据表', '源数据表', '源表名', '源表别名', '表名', '数据源表']:
                    if row_dict.get(key):
                        source_table = str(row_dict.get(key, ''))
                        break
                
                for key in ['源字段', '源字段名(英文)', '字段名(英文)', '源字段英文名', '字段名', '源字段名']:
                    if row_dict.get(key):
                        source_field = str(row_dict.get(key, ''))
                        break
                
                for key in ['源字段名(中文)', '字段名(中文)', '源字段中文名', '字段中文名']:
                    if row_dict.get(key):
                        source_field_cn = str(row_dict.get(key, ''))
                        break
                
                field_mappings.append({
                    'row_index': row_idx,
                    'group_id': str(row_dict.get('组别编号', '')),
                    'field_seq': row_dict.get('字段序号', 0),
                    'target_field': str(target_field).strip(),
                    'target_field_cn': str(row_dict.get('字段中文名称', '') or 
                                          row_dict.get('目标字段中文名', '')),
                    'source_table': source_table,
                    'source_table_alias': source_table,  # 同时存储为source_table_alias
                    'source_field': source_field,
                    'source_field_cn': source_field_cn,
                    'data_type': str(row_dict.get('数据类型', '')),
                    'default_value': str(row_dict.get('默认值', '')),
                    'transformation_logic': str(row_dict.get('字段加工逻辑', '') or 
                                               row_dict.get('转换逻辑', '')),
                    'extract_method': str(row_dict.get('取数方式', '')),
                    'raw_data': row_dict
                })
        
        return field_mappings
    
    def _update_progress(self, task_id: str, phase: str, step: str, 
                        current: int, total: int):
        if self.progress_callback:
            progress = ProgressInfo(
                current_phase=phase,
                current_step=step,
                total_items=total,
                processed_items=current,
                percentage=(current / total * 100) if total > 0 else 0,
                start_time=datetime.now()
            )
            self.progress_callback(task_id, progress.to_dict())


class SQLProcessor:
    """SQL处理器 - 解析SQL存储过程并提取结构信息"""
    
    def __init__(self, progress_callback: Optional[Callable] = None):
        self.progress_callback = progress_callback
        self.logger = logging.getLogger(__name__)
        self.retry_mechanism = RetryMechanism(max_retries=2)
    
    def parse_sql_file(self, file_path: str, 
                      task_id: str) -> Tuple[bool, Dict, Optional[ErrorInfo]]:
        try:
            self._update_progress(task_id, "parse", "读取SQL文件", 0, 100)
            
            def _read_file():
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            
            sql_content = self.retry_mechanism.execute_with_retry(
                _read_file,
                retryable_exceptions=(IOError, OSError, UnicodeDecodeError)
            )
            
            self._update_progress(task_id, "parse", "预处理SQL内容", 25, 100)
            preprocessed = self._preprocess_sql(sql_content)
            
            self._update_progress(task_id, "parse", "提取INSERT块", 50, 100)
            insert_blocks = self._extract_insert_blocks(preprocessed)
            
            self._update_progress(task_id, "parse", "解析字段映射", 75, 100)
            result = {
                'file_path': file_path,
                'file_hash': hashlib.md5(sql_content.encode()).hexdigest(),
                'content_length': len(sql_content),
                'procedure_name': self._extract_procedure_name(sql_content),
                'insert_blocks': insert_blocks,
                'parse_metadata': {
                    'parse_time': datetime.now().isoformat(),
                    'total_blocks': len(insert_blocks),
                    'total_lines': len(sql_content.split('\n'))
                }
            }
            
            self._update_progress(task_id, "parse", "解析完成", 100, 100)
            
            self.logger.info(f"SQL解析完成: {file_path}, "
                           f"INSERT块={len(insert_blocks)}")
            
            return True, result, None
            
        except Exception as e:
            error_info = ErrorInfo(
                error_type=type(e).__name__,
                error_message=str(e),
                error_phase="parse",
                error_step="sql_file",
                timestamp=datetime.now(),
                stack_trace=traceback.format_exc(),
                is_recoverable=True
            )
            self.logger.error(f"SQL解析失败: {e}")
            return False, {}, error_info
    
    def _preprocess_sql(self, sql_content: str) -> str:
        lines = sql_content.split('\n')
        cleaned_lines = []
        
        for line in lines:
            if line.strip().startswith('--'):
                continue
            if '/*' in line and '*/' in line:
                line = re.sub(r'/\*.*?\*/', '', line)
            cleaned_lines.append(line)
        
        return '\n'.join(cleaned_lines)
    
    def _extract_insert_blocks(self, sql_content: str) -> List[Dict]:
        blocks = []
        
        insert_pattern = re.compile(
            r'INSERT\s+INTO\s+(\w+)(?:\s+PARTITION\s*\(([^)]+)\))?\s*\(([\s\S]*?)\)\s*SELECT\s+([\s\S]*?)\s+FROM\s+([\s\S]*?)(?:WHERE\s+([\s\S]*?))?(?:;|$)',
            re.IGNORECASE
        )
        
        for match in insert_pattern.finditer(sql_content):
            target_table = match.group(1)
            partition = match.group(2) or ''
            insert_fields_str = match.group(3)
            select_values_str = match.group(4)
            from_clause_str = match.group(5)
            where_clause_str = match.group(6) or ''
            
            insert_fields = self._parse_field_list(insert_fields_str)
            select_values = self._parse_select_values(select_values_str)
            from_clause, join_clauses = self._parse_from_clause(from_clause_str)
            
            field_mapping = {}
            for i, field in enumerate(insert_fields):
                if i < len(select_values):
                    field_mapping[field] = select_values[i]
            
            group_id = self._extract_group_id(sql_content, match.start())
            
            blocks.append({
                'target_table': target_table,
                'partition': partition,
                'group_id': group_id,
                'insert_fields': insert_fields,
                'select_values': select_values,
                'field_mapping': field_mapping,
                'from_clause': from_clause,
                'join_clauses': join_clauses,
                'where_clause': where_clause_str.strip(),
                'raw_sql': match.group(0)[:1000]
            })
        
        return blocks
    
    def _parse_field_list(self, fields_str: str) -> List[str]:
        fields = []
        for line in fields_str.split('\n'):
            clean_line = re.sub(r'--.*$', '', line).strip()
            if clean_line:
                field = clean_line.rstrip(',').strip()
                if field:
                    fields.append(field)
        return fields
    
    def _parse_select_values(self, select_str: str) -> List[str]:
        values = []
        current = ''
        paren_depth = 0
        case_depth = 0
        
        for char in select_str:
            if char == '(':
                paren_depth += 1
                current += char
            elif char == ')':
                paren_depth -= 1
                current += char
            elif char == ',' and paren_depth == 0 and case_depth == 0:
                if current.strip():
                    values.append(current.strip())
                current = ''
            else:
                current += char
            
            if current.upper().endswith('CASE'):
                case_depth += 1
            elif case_depth > 0 and current.upper().rstrip().endswith('END'):
                case_depth -= 1
        
        if current.strip():
            values.append(current.strip())
        
        return values
    
    def _parse_from_clause(self, from_str: str) -> Tuple[str, List[Dict]]:
        from_str = from_str.strip()
        
        from_match = re.search(r'^(\w+(?:\s+\w+)?)', from_str, re.IGNORECASE)
        from_clause = from_match.group(1) if from_match else ''
        
        join_clauses = []
        join_pattern = re.compile(
            r'(LEFT\s+JOIN|RIGHT\s+JOIN|INNER\s+JOIN|JOIN)\s+(\w+(?:\s+\w+)?)\s+ON\s+([\s\S]*?)(?=(?:LEFT|RIGHT|INNER|JOIN|WHERE|;|$))',
            re.IGNORECASE
        )
        
        for match in join_pattern.finditer(from_str):
            join_clauses.append({
                'type': match.group(1).upper(),
                'table': match.group(2).strip(),
                'condition': match.group(3).strip().rstrip('AND').strip()
            })
        
        return from_clause, join_clauses
    
    def _extract_group_id(self, sql_content: str, position: int) -> str:
        before_sql = sql_content[:position]
        lines = before_sql.split('\n')[-20:]
        
        for line in reversed(lines):
            match = re.search(r'(MP\d+)', line, re.IGNORECASE)
            if match:
                return match.group(1).upper()
        
        return ''
    
    def _extract_procedure_name(self, sql_content: str) -> str:
        match = re.search(r'CREATE\s+OR\s+REPLACE\s+PROCEDURE\s+(\w+)', 
                         sql_content, re.IGNORECASE)
        return match.group(1) if match else ''
    
    def _update_progress(self, task_id: str, phase: str, step: str, 
                        current: int, total: int):
        if self.progress_callback:
            progress = ProgressInfo(
                current_phase=phase,
                current_step=step,
                total_items=total,
                processed_items=current,
                percentage=(current / total * 100) if total > 0 else 0,
                start_time=datetime.now()
            )
            self.progress_callback(task_id, progress.to_dict())


class DatabaseStorageService:
    """数据库存储服务 - 存储解析后的数据到数据库"""
    
    def __init__(self, db_service, progress_callback: Optional[Callable] = None):
        self.db_service = db_service
        self.progress_callback = progress_callback
        self.logger = logging.getLogger(__name__)
        self.retry_mechanism = RetryMechanism(max_retries=3)
    
    def store_mapping_data(self, task_id: str, mapping_data: Dict) -> Tuple[bool, Optional[ErrorInfo]]:
        try:
            self._update_progress(task_id, "store", "存储数据源信息", 0, 100)
            
            data_sources = mapping_data.get('data_sources', [])
            field_mappings = mapping_data.get('field_mappings', [])
            
            total_items = len(data_sources) + len(field_mappings)
            processed = 0
            
            batch_data_sources = []
            dataset_comparison_results = []
            for ds in data_sources:
                batch_data_sources.append({
                    'mapping_file': mapping_data.get('file_path', ''),
                    'sheet_name': ds.get('sheet_name', ''),
                    'alias': ds.get('table_alias', ''),
                    'name': ds.get('table_name', ''),
                    'group_id': ds.get('group_id', '')
                })
                
                # 同时存储到 dataset_comparison_results 表
                dataset_comparison_results.append({
                    'task_id': task_id,
                    'group_id': ds.get('group_id', ''),
                    'source_tables': f"{ds.get('table_alias', '')}:{ds.get('table_name', '')}",
                    'compare_module': '数据源',
                    'mapping_item': ds.get('table_name', ''),
                    'sql_content': '',
                    'is_consistent': '',
                    'remark': ''
                })
                
                processed += 1
                self._update_progress(task_id, "store", "存储数据源信息", 
                                     processed, total_items)
            
            if batch_data_sources:
                self._batch_insert_data_sources(batch_data_sources)
            
            # 存储到 dataset_comparison_results 表
            if dataset_comparison_results:
                self.db_service.save_dataset_comparison_batch(task_id, dataset_comparison_results)
            
            batch_field_mappings = []
            datamap_comparison_results = []
            for fm in field_mappings:
                batch_field_mappings.append({
                    'mapping_file': mapping_data.get('file_path', ''),
                    'group_id': fm.get('group_id', ''),
                    'field_seq': fm.get('field_seq', 0),
                    'target_field': fm.get('target_field', ''),
                    'target_field_cn': fm.get('target_field_cn', ''),
                    'source_table_alias': fm.get('source_table', ''),
                    'source_field': fm.get('source_field', ''),
                    'source_field_cn': fm.get('source_field_cn', ''),
                    'data_type': fm.get('data_type', ''),
                    'default_value': fm.get('default_value', ''),
                    'transformation_logic': fm.get('transformation_logic', ''),
                    'extract_method': fm.get('extract_method', '')
                })
                
                # 同时存储到 datamap_comparison_results 表
                datamap_comparison_results.append({
                    'task_id': task_id,
                    'group_id': fm.get('group_id', ''),
                    'field_seq': fm.get('field_seq', 0),
                    'target_field_cn': fm.get('target_field_cn', ''),
                    'target_field_en': fm.get('target_field', ''),
                    'extract_method': fm.get('extract_method', ''),
                    'source_table': fm.get('source_table_alias', '') or fm.get('source_table', ''),
                    'source_field_en': fm.get('source_field', ''),
                    'source_field_cn': fm.get('source_field_cn', ''),
                    'default_value': fm.get('default_value', ''),
                    'transform_logic': fm.get('transformation_logic', ''),
                    'sql_expression': '',
                    'is_consistent': '',
                    'remark': ''
                })
                
                processed += 1
                self._update_progress(task_id, "store", "存储字段映射", 
                                     processed, total_items)
            
            if batch_field_mappings:
                self._batch_insert_field_mappings(batch_field_mappings)
            
            # 存储到 datamap_comparison_results 表
            if datamap_comparison_results:
                self.db_service.save_datamap_comparison_batch(task_id, datamap_comparison_results)
            
            self.logger.info(f"数据存储完成: task_id={task_id}, "
                           f"数据源={len(data_sources)}, 字段映射={len(field_mappings)}")
            
            return True, None
            
        except Exception as e:
            error_info = ErrorInfo(
                error_type=type(e).__name__,
                error_message=str(e),
                error_phase="store",
                error_step="mapping_data",
                timestamp=datetime.now(),
                stack_trace=traceback.format_exc(),
                is_recoverable=True
            )
            self.logger.error(f"存储映射数据失败: {e}")
            return False, error_info
    
    def store_sql_data(self, task_id: str, sql_data: Dict) -> Tuple[bool, Optional[ErrorInfo]]:
        try:
            self._update_progress(task_id, "store", "存储SQL文件信息", 0, 100)
            
            self.db_service.save_sql_file(
                sql_id=task_id,
                filename=os.path.basename(sql_data.get('file_path', '')),
                content=sql_data.get('raw_content', '')
            )
            
            insert_blocks = sql_data.get('insert_blocks', [])
            total_blocks = len(insert_blocks)
            
            # 准备更新 dataset_comparison_results 和 datamap_comparison_results 表的 SQL 信息
            sql_field_mappings = {}
            for block in insert_blocks:
                group_id = block.get('group_id', '')
                if group_id not in sql_field_mappings:
                    sql_field_mappings[group_id] = {}
                sql_field_mappings[group_id].update(block.get('field_mapping', {}))
            
            # 更新 dataset_comparison_results 表中的 SQL 信息
            for block in insert_blocks:
                group_id = block.get('group_id', '')
                from_clause = block.get('from_clause', '')
                join_clauses = block.get('join_clauses', [])
                where_clause = block.get('where_clause', '')
                
                # 构建 SQL 内容
                sql_content = f"FROM {from_clause}"
                for join in join_clauses:
                    sql_content += f"\n{join.get('type', '')} {join.get('table', '')} ON {join.get('condition', '')}"
                if where_clause:
                    sql_content += f"\nWHERE {where_clause}"
                
                # 更新 dataset_comparison_results 表
                self.db_service.update_dataset_comparison_sql(task_id, group_id, sql_content)
            
            # 更新 datamap_comparison_results 表中的 SQL 信息
            for group_id, field_mappings in sql_field_mappings.items():
                for target_field, sql_expression in field_mappings.items():
                    self.db_service.update_datamap_comparison_sql(task_id, group_id, target_field, sql_expression)
            
            for idx, block in enumerate(insert_blocks):
                self._update_progress(task_id, "store", f"存储INSERT块 {idx + 1}/{total_blocks}",
                                     idx + 1, total_blocks)
            
            self.logger.info(f"SQL数据存储完成: task_id={task_id}, "
                           f"INSERT块={len(insert_blocks)}")
            
            return True, None
            
        except Exception as e:
            error_info = ErrorInfo(
                error_type=type(e).__name__,
                error_message=str(e),
                error_phase="store",
                error_step="sql_data",
                timestamp=datetime.now(),
                stack_trace=traceback.format_exc(),
                is_recoverable=True
            )
            self.logger.error(f"存储SQL数据失败: {e}")
            return False, error_info
    
    def _batch_insert_data_sources(self, data_sources: List[Dict]):
        for ds in data_sources:
            self.db_service.save_data_source(
                mapping_file=ds.get('mapping_file', ''),
                sheet_name=ds.get('sheet_name', ''),
                alias=ds.get('alias', ''),
                name=ds.get('name', ''),
                relationships=None
            )
    
    def _batch_insert_field_mappings(self, field_mappings: List[Dict]):
        for fm in field_mappings:
            self.db_service.save_field_mapping(
                mapping_file=fm.get('mapping_file', ''),
                group_id=fm.get('group_id', ''),
                field_seq=fm.get('field_seq', 0),
                target_field=fm.get('target_field', ''),
                target_field_cn=fm.get('target_field_cn', ''),
                source_table_alias=fm.get('source_table_alias', ''),
                source_field=fm.get('source_field', ''),
                source_field_cn=fm.get('source_field_cn', ''),
                default_value=fm.get('default_value', ''),
                transformation_logic=fm.get('transformation_logic', ''),
                data_type=fm.get('data_type', ''),
                extract_method=fm.get('extract_method', ''),
                sheet=fm.get('sheet', '')
            )
    
    def _update_progress(self, task_id: str, phase: str, step: str, 
                        current: int, total: int):
        if self.progress_callback:
            progress = ProgressInfo(
                current_phase=phase,
                current_step=step,
                total_items=total,
                processed_items=current,
                percentage=(current / total * 100) if total > 0 else 0,
                start_time=datetime.now()
            )
            self.progress_callback(task_id, progress.to_dict())


class ConsistencyComparator:
    """一致性对比器 - 执行全面的数据一致性对比"""
    
    def __init__(self, db_service, logger_instance: Optional[ComparisonLogger] = None,
                 progress_callback: Optional[Callable] = None):
        self.db_service = db_service
        self.comparison_logger = logger_instance or ComparisonLogger()
        self.progress_callback = progress_callback
        self.logger = logging.getLogger(__name__)
        self.retry_mechanism = RetryMechanism(max_retries=3)
    
    def compare_all_records(self, task_id: str, mapping_file: str, 
                           sql_data: Dict) -> Tuple[bool, ComparisonMetrics, Optional[ErrorInfo]]:
        try:
            field_mappings = self.db_service.get_field_mappings(mapping_file)
            
            if not field_mappings:
                return True, ComparisonMetrics(), None
            
            total_records = len(field_mappings)
            self.comparison_logger.start_comparison(total_records)
            
            insert_blocks = sql_data.get('insert_blocks', [])
            sql_field_mapping = self._build_sql_field_mapping(insert_blocks)
            
            for idx, mapping_record in enumerate(field_mappings):
                self._update_progress(task_id, "compare", 
                                     f"对比记录 {idx + 1}/{total_records}",
                                     idx + 1, total_records)
                
                start_time = time.time()
                
                try:
                    result = self._compare_single_record(
                        mapping_record, sql_field_mapping, insert_blocks
                    )
                    duration = time.time() - start_time
                    
                    self.comparison_logger.log_record_comparison(
                        record_id=f"{mapping_record.get('group_id', '')}_{mapping_record.get('field_seq', '')}",
                        status=result['status'],
                        details=result['details'],
                        duration=duration
                    )
                    
                    self._save_comparison_result(task_id, mapping_record, result)
                    
                except Exception as e:
                    self.comparison_logger.log_record_comparison(
                        record_id=f"{mapping_record.get('group_id', '')}_{mapping_record.get('field_seq', '')}",
                        status='error',
                        details={'error': str(e)},
                        duration=time.time() - start_time
                    )
            
            metrics = self.comparison_logger.end_comparison()
            
            self.logger.info(f"对比完成: task_id={task_id}, {metrics.to_dict()}")
            
            return True, metrics, None
            
        except Exception as e:
            error_info = ErrorInfo(
                error_type=type(e).__name__,
                error_message=str(e),
                error_phase="compare",
                error_step="compare_all",
                timestamp=datetime.now(),
                stack_trace=traceback.format_exc(),
                is_recoverable=True
            )
            self.logger.error(f"一致性对比失败: {e}")
            return False, ComparisonMetrics(), error_info
    
    def _compare_single_record(self, mapping_record: Dict, 
                               sql_field_mapping: Dict,
                               insert_blocks: List[Dict]) -> Dict:
        result = {
            'status': 'consistent',
            'details': {},
            'discrepancies': []
        }
        
        group_id = mapping_record.get('group_id', '')
        target_field = mapping_record.get('target_field', '')
        source_field = mapping_record.get('source_field', '')
        transformation_logic = mapping_record.get('transformation_logic', '')
        default_value = mapping_record.get('default_value', '')
        
        sql_expression = sql_field_mapping.get(group_id, {}).get(target_field, '')
        
        if not sql_expression:
            result['status'] = 'warning'
            result['details']['sql_expression'] = '未在SQL中找到对应字段'
            result['discrepancies'].append({
                'type': 'missing_field',
                'field': target_field,
                'severity': 'high',
                'message': f'字段 {target_field} 在SQL中未找到'
            })
            return result
        
        result['details']['sql_expression'] = sql_expression
        
        if source_field and source_field not in sql_expression:
            if not self._is_field_in_expression(source_field, sql_expression):
                result['status'] = 'inconsistent'
                result['discrepancies'].append({
                    'type': 'source_field_mismatch',
                    'field': target_field,
                    'mapping_value': source_field,
                    'sql_value': sql_expression,
                    'severity': 'high',
                    'message': f'源字段不匹配: Mapping={source_field}, SQL={sql_expression}'
                })
        
        if transformation_logic and transformation_logic.strip():
            if not self._check_transformation_logic(transformation_logic, sql_expression):
                result['status'] = 'inconsistent'
                result['discrepancies'].append({
                    'type': 'transformation_mismatch',
                    'field': target_field,
                    'mapping_value': transformation_logic,
                    'sql_value': sql_expression,
                    'severity': 'medium',
                    'message': f'转换逻辑不匹配: Mapping={transformation_logic}'
                })
        
        if default_value and default_value.strip():
            if not self._check_default_value(default_value, sql_expression):
                result['status'] = 'inconsistent'
                result['discrepancies'].append({
                    'type': 'default_value_mismatch',
                    'field': target_field,
                    'mapping_value': default_value,
                    'sql_value': sql_expression,
                    'severity': 'medium',
                    'message': f'默认值不匹配: Mapping={default_value}'
                })
        
        return result
    
    def _build_sql_field_mapping(self, insert_blocks: List[Dict]) -> Dict:
        mapping = {}
        for block in insert_blocks:
            group_id = block.get('group_id', '')
            if group_id not in mapping:
                mapping[group_id] = {}
            mapping[group_id].update(block.get('field_mapping', {}))
        return mapping
    
    def _is_field_in_expression(self, field: str, expression: str) -> bool:
        patterns = [
            rf'\b{re.escape(field)}\b',
            rf'\.{re.escape(field)}\b',
            rf'\b{re.escape(field)}\s*,',
            rf'\b{re.escape(field)}\s+AS',
        ]
        return any(re.search(p, expression, re.IGNORECASE) for p in patterns)
    
    def _check_transformation_logic(self, logic: str, expression: str) -> bool:
        logic_lower = logic.lower()
        expression_lower = expression.lower()
        
        if 'substr' in logic_lower or '截取' in logic_lower:
            if 'substr' not in expression_lower and 'substring' not in expression_lower:
                return False
        
        if 'nvl' in logic_lower or 'coalesce' in logic_lower:
            if 'nvl' not in expression_lower and 'coalesce' not in expression_lower:
                return False
        
        if 'case' in logic_lower:
            if 'case' not in expression_lower:
                return False
        
        return True
    
    def _check_default_value(self, default: str, expression: str) -> bool:
        default_clean = default.strip().strip("'\"")
        return default_clean in expression or f"'{default_clean}'" in expression
    
    def _save_comparison_result(self, task_id: str, mapping_record: Dict, result: Dict):
        is_consistent = '是' if result['status'] == 'consistent' else '否'
        remark = '; '.join([d['message'] for d in result.get('discrepancies', [])])
        
        self.db_service.save_datamap_comparison(
            task_id=task_id,
            group_id=mapping_record.get('group_id', ''),
            field_seq=mapping_record.get('field_seq', 0),
            target_field_cn=mapping_record.get('target_field_cn', ''),
            target_field_en=mapping_record.get('target_field', ''),
            extract_method=mapping_record.get('extract_method', ''),
            source_table=mapping_record.get('source_table_alias', ''),
            source_field_en=mapping_record.get('source_field', ''),
            source_field_cn=mapping_record.get('source_field_cn', ''),
            default_value=mapping_record.get('default_value', ''),
            transform_logic=mapping_record.get('transformation_logic', ''),
            sql_expression=result['details'].get('sql_expression', ''),
            is_consistent=is_consistent,
            remark=remark
        )
    
    def _update_progress(self, task_id: str, phase: str, step: str, 
                        current: int, total: int):
        if self.progress_callback:
            progress = ProgressInfo(
                current_phase=phase,
                current_step=step,
                total_items=total,
                processed_items=current,
                percentage=(current / total * 100) if total > 0 else 0,
                start_time=datetime.now()
            )
            self.progress_callback(task_id, progress.to_dict())


class DataFlowManager:
    """数据流管理器 - 协调整个数据流处理过程"""
    
    def __init__(self, upload_dir: str, db_service, 
                 progress_callback: Optional[Callable] = None):
        self.upload_dir = upload_dir
        self.db_service = db_service
        self.progress_callback = progress_callback
        self.logger = logging.getLogger(__name__)
        
        self.upload_handler = DataUploadHandler(upload_dir, self.logger)
        self.mapping_processor = ITMappingProcessor(self._internal_progress_callback)
        self.sql_processor = SQLProcessor(self._internal_progress_callback)
        self.storage_service = DatabaseStorageService(db_service, self._internal_progress_callback)
        self.comparison_logger = ComparisonLogger()
        self.comparator = ConsistencyComparator(
            db_service, self.comparison_logger, self._internal_progress_callback
        )
        
        self._tasks: Dict[str, Dict] = {}
        self._lock = threading.Lock()
    
    def _internal_progress_callback(self, task_id: str, progress: Dict):
        with self._lock:
            if task_id in self._tasks:
                self._tasks[task_id]['progress'] = progress
                self._tasks[task_id]['progress_info'] = progress
        
        if self.progress_callback:
            self.progress_callback(task_id, progress)
    
    def start_data_flow(self, mapping_file_content: bytes, mapping_filename: str,
                       sql_file_content: bytes, sql_filename: str) -> Dict:
        task_id = str(uuid.uuid4())
        
        with self._lock:
            self._tasks[task_id] = {
                'status': DataFlowStatus.PENDING.value,
                'progress': ProgressInfo().to_dict(),
                'progress_info': {},
                'errors': [],
                'result': {},
                'start_time': datetime.now()
            }
        
        thread = threading.Thread(
            target=self._execute_data_flow,
            args=(task_id, mapping_file_content, mapping_filename, 
                  sql_file_content, sql_filename)
        )
        thread.daemon = True
        thread.start()
        
        return {
            'success': True,
            'data': {
                'task_id': task_id,
                'status': 'started',
                'message': '数据流处理已启动'
            }
        }
    
    def _execute_data_flow(self, task_id: str, mapping_file_content: bytes, 
                          mapping_filename: str, sql_file_content: bytes, 
                          sql_filename: str):
        try:
            self._update_task_status(task_id, DataFlowStatus.IN_PROGRESS)
            
            success, mapping_path, error = self.upload_handler.save_uploaded_file(
                mapping_file_content, mapping_filename, 'mapping'
            )
            if not success:
                self._add_task_error(task_id, error)
                self._update_task_status(task_id, DataFlowStatus.FAILED)
                return
            
            success, sql_path, error = self.upload_handler.save_uploaded_file(
                sql_file_content, sql_filename, 'sql'
            )
            if not success:
                self._add_task_error(task_id, error)
                self._update_task_status(task_id, DataFlowStatus.FAILED)
                return
            
            sql_content = sql_file_content.decode('utf-8') if isinstance(sql_file_content, bytes) else sql_file_content
            
            self._internal_progress_callback(task_id, {
                'current_phase': 'step1',
                'current_step': '解析IT-Mapping并存入数据库',
                'total_items': 100,
                'processed_items': 0,
                'percentage': 0,
                'message': '第一步：解析IT-Mapping文件并存入数据库'
            })
            
            from extract_datamap_to_md import extract_and_store_to_db
            step1_result = extract_and_store_to_db(
                excel_path=mapping_path,
                task_id=task_id,
                db_path=self.db_service.db_path
            )
            
            if not step1_result.get('success'):
                error_info = ErrorInfo(
                    error_type="Step1Error",
                    error_message="第一步执行失败：IT-Mapping解析或数据库存储失败",
                    error_phase="step1",
                    error_step="extract_and_store_to_db",
                    timestamp=datetime.now(),
                    is_recoverable=False
                )
                self._add_task_error(task_id, error_info)
                self._update_task_status(task_id, DataFlowStatus.FAILED)
                return
            
            self._internal_progress_callback(task_id, {
                'current_phase': 'step2',
                'current_step': '提取SQL表达式',
                'total_items': 100,
                'processed_items': 50,
                'percentage': 50,
                'message': '第二步：从SQL文件中提取对应的SQL表达式'
            })
            
            self._internal_progress_callback(task_id, {
                'current_phase': 'step3',
                'current_step': '一致性对比',
                'total_items': 100,
                'processed_items': 75,
                'percentage': 75,
                'message': '第三步：对比IT-Mapping与存储过程SQL，更新数据库'
            })
            
            from sql_extraction_agent import run_step2_and_step3
            import json
            import os
            
            ollama_config = None
            config_path = os.path.join(os.path.dirname(__file__), '..', 'config', 'ollama_config.json')
            if os.path.exists(config_path):
                try:
                    with open(config_path, 'r', encoding='utf-8') as f:
                        ollama_config = json.load(f)
                except Exception as e:
                    self.logger.warning(f"加载Ollama配置失败: {e}")
            
            step2_3_result = run_step2_and_step3(
                task_id=task_id,
                sql_content=sql_content,
                db_path=self.db_service.db_path,
                enable_agent_review=True,
                ollama_config=ollama_config
            )
            
            if not step2_3_result.get('success'):
                error_info = ErrorInfo(
                    error_type="Step2_3Error",
                    error_message="第二步或第三步执行失败",
                    error_phase="step2_3",
                    error_step="run_step2_and_step3",
                    timestamp=datetime.now(),
                    is_recoverable=False
                )
                self._add_task_error(task_id, error_info)
                self._update_task_status(task_id, DataFlowStatus.FAILED)
                return
            
            datamap_result = step2_3_result.get('datamap', {})
            dataset_result = step2_3_result.get('dataset', {})
            
            review_results = datamap_result.get('review_results', {})
            if datamap_result.get('agent_reviewed') and review_results:
                consistent_after_review = datamap_result.get('consistent', 0) + review_results.get('overturned_count', 0)
                inconsistent_after_review = datamap_result.get('inconsistent', 0) - review_results.get('overturned_count', 0)
            else:
                consistent_after_review = datamap_result.get('consistent', 0)
                inconsistent_after_review = datamap_result.get('inconsistent', 0)
            
            with self._lock:
                self._tasks[task_id]['result'] = {
                    'mapping_file': mapping_filename,
                    'sql_file': sql_filename,
                    'metrics': {
                        'total_records': datamap_result.get('total', 0),
                        'consistent_records': consistent_after_review + dataset_result.get('consistent', 0),
                        'inconsistent_records': inconsistent_after_review + dataset_result.get('inconsistent', 0),
                        'warning_records': 0,
                        'error_records': datamap_result.get('errors', 0) + dataset_result.get('errors', 0),
                        'skipped_records': 0,
                        'comparison_time': 0,
                        'avg_comparison_time': 0,
                        'agent_reviewed': datamap_result.get('agent_reviewed', False),
                        'overturned_count': review_results.get('overturned_count', 0) if review_results else 0
                    },
                    'mapping_path': mapping_path,
                    'sql_path': sql_path,
                    'step1_result': step1_result,
                    'step2_3_result': step2_3_result
                }
            
            self._internal_progress_callback(task_id, {
                'current_phase': 'complete',
                'current_step': '处理完成',
                'total_items': 100,
                'processed_items': 100,
                'percentage': 100,
                'message': '三步流程全部完成'
            })
            
            self._update_task_status(task_id, DataFlowStatus.COMPLETED)
            
            self._save_to_history(task_id, mapping_filename, sql_filename, sql_content)
            
        except Exception as e:
            error_info = ErrorInfo(
                error_type=type(e).__name__,
                error_message=str(e),
                error_phase="data_flow",
                error_step="execute",
                timestamp=datetime.now(),
                stack_trace=traceback.format_exc(),
                is_recoverable=False
            )
            self._add_task_error(task_id, error_info)
            self._update_task_status(task_id, DataFlowStatus.FAILED)
    
    def execute_step2_and_step3_only(self, task_id: str, sql_content: str) -> Dict:
        """仅执行第二步和第三步
        
        用于已经完成第一步后，单独执行SQL提取和对比
        
        Args:
            task_id: 任务ID
            sql_content: SQL文件内容
            
        Returns:
            Dict: 处理结果
        """
        try:
            self._internal_progress_callback(task_id, {
                'current_phase': 'step2',
                'current_step': '提取SQL表达式',
                'total_items': 100,
                'processed_items': 0,
                'percentage': 0,
                'message': '第二步：从SQL文件中提取对应的SQL表达式'
            })
            
            self._internal_progress_callback(task_id, {
                'current_phase': 'step3',
                'current_step': '一致性对比',
                'total_items': 100,
                'processed_items': 50,
                'percentage': 50,
                'message': '第三步：对比IT-Mapping与存储过程SQL，更新数据库'
            })
            
            from sql_extraction_agent import run_step2_and_step3
            result = run_step2_and_step3(
                task_id=task_id,
                sql_content=sql_content,
                db_path=self.db_service.db_path
            )
            
            self._internal_progress_callback(task_id, {
                'current_phase': 'complete',
                'current_step': '处理完成',
                'total_items': 100,
                'processed_items': 100,
                'percentage': 100,
                'message': '第二步和第三步完成'
            })
            
            return result
            
        except Exception as e:
            self.logger.error(f"执行第二步和第三步失败: {e}")
            return {
                'success': False,
                'error': str(e)
            }
    
    def get_task_status(self, task_id: str) -> Optional[Dict]:
        with self._lock:
            if task_id in self._tasks:
                task = self._tasks[task_id]
                return {
                    'task_id': task_id,
                    'status': task['status'],
                    'progress': task['progress'],
                    'errors': [e.to_dict() for e in task['errors']],
                    'result': task.get('result', {}),
                    'start_time': task['start_time'].isoformat() if task.get('start_time') else None
                }
        return None
    
    def get_task_progress(self, task_id: str) -> Dict:
        """获取任务进度"""
        with self._lock:
            if task_id in self._tasks:
                task = self._tasks[task_id]
                percentage = task['progress_info'].get('percentage', 0)
                return {
                    'success': True,
                    'data': {
                        'task_id': task_id,
                        'status': task['status'],
                        'progress': percentage,
                        'current_step': task['progress_info'].get('current_step', ''),
                        'percentage': percentage,
                        'message': task['progress_info'].get('message', '')
                    }
                }
        return {'success': False, 'status': 'not_found', 'message': f'任务 {task_id} 不存在'}
    
    def get_task_result(self, task_id: str) -> Dict:
        """获取任务结果"""
        with self._lock:
            if task_id in self._tasks:
                task = self._tasks[task_id]
                if task['status'] == 'completed':
                    result = task.get('result', {})
                    metrics = result.get('metrics', {})
                    return {
                        'success': True,
                        'data': {
                            'task_id': task_id,
                            'status': 'completed',
                            'mapping_file': result.get('mapping_file', ''),
                            'sql_filename': result.get('sql_file', ''),
                            'total_checks': metrics.get('total_records', 0),
                            'consistent_count': metrics.get('consistent_records', 0),
                            'inconsistent_count': metrics.get('inconsistent_records', 0),
                            'agent_reviewed': metrics.get('agent_reviewed', False),
                            'overturned_count': metrics.get('overturned_count', 0),
                            'findings': []
                        }
                    }
                elif task['status'] == 'failed':
                    return {
                        'success': False,
                        'data': {
                            'task_id': task_id,
                            'status': 'failed',
                            'errors': [e.to_dict() for e in task['errors']]
                        }
                    }
                else:
                    return {
                        'success': True,
                        'data': {
                            'task_id': task_id,
                            'status': task['status'],
                            'progress': task['progress']
                        }
                    }
        return {'success': False, 'message': f'任务 {task_id} 不存在'}
    
    def _update_task_status(self, task_id: str, status: DataFlowStatus):
        with self._lock:
            if task_id in self._tasks:
                self._tasks[task_id]['status'] = status.value
                self.logger.info(f"任务 {task_id} 状态更新: {status.value}")
    
    def _add_task_error(self, task_id: str, error: ErrorInfo):
        with self._lock:
            if task_id in self._tasks:
                self._tasks[task_id]['errors'].append(error)
                self.logger.error(f"任务 {task_id} 错误: {error.error_message}")
    
    def _save_to_history(self, task_id: str, mapping_filename: str, sql_filename: str, sql_content: str):
        """保存任务结果到历史记录"""
        try:
            from services.history_service import HistoryService
            history_service = HistoryService()
            
            with self._lock:
                if task_id not in self._tasks:
                    return
                task = self._tasks[task_id]
                result = task.get('result', {})
                metrics = result.get('metrics', {})
            
            history_data = {
                'task_id': task_id,
                'mapping_file': mapping_filename,
                'sql_filename': sql_filename,
                'sql_content': sql_content,
                'status': 'completed',
                'total_checks': metrics.get('total_records', 0),
                'consistent_count': metrics.get('consistent_records', 0),
                'inconsistent_count': metrics.get('inconsistent_records', 0),
                'duration': 0,
                'findings': [],
                'summary_text': f"规则引擎预检 + Agent复核完成。一致: {metrics.get('consistent_records', 0)}，不一致: {metrics.get('inconsistent_records', 0)}"
            }
            
            history_service.save_review_result(task_id, history_data)
            self.logger.info(f"任务 {task_id} 已保存到历史记录")
            
        except Exception as e:
            self.logger.error(f"保存任务 {task_id} 到历史记录失败: {e}")
    
    def get_comparison_log(self, task_id: str) -> Optional[Dict]:
        return {
            'metrics': self.comparison_logger.metrics.to_dict(),
            'log_entries': self.comparison_logger.comparison_log
        }
    
    def export_comparison_report(self, task_id: str, output_path: str) -> bool:
        return self.comparison_logger.export_log(output_path)
